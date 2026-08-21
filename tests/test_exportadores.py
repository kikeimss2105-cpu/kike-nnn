from io import BytesIO

import pandas as pd
from docx import Document
from openpyxl import load_workbook

from engine.docente import analizar_sesion
from utils.exportadores import generar_excel, generar_word, generar_word_docente


def _plan_minimo() -> pd.DataFrame:
    return pd.DataFrame([{
        "Código": "00030",
        "NANDA": "Deterioro del intercambio gaseoso",
        "Puntaje": 8,
        "Confianza": "Media",
        "Jerarquía": "Complementario",
        "Prioridad": "Alta",
        "NOC sugerido": "Estado respiratorio: intercambio gaseoso",
        "NIC sugerido": "Monitorización respiratoria",
        "Meta esperada": "Meta educativa de prueba",
        "Indicadores NOC": "Indicador de prueba",
        "Actividades NIC": "Actividad de prueba",
        "Fundamentos": "Fundamento de prueba",
        "Nota": "Requiere validación clínica",
    }])


def test_excel_generado_es_abrible_y_conserva_advertencia() -> None:
    salida = generar_excel(_plan_minimo(), {"Tipo de paciente": "Adulto"})
    libro = load_workbook(BytesIO(salida.getvalue()), read_only=True)

    assert "Plan NNN completo" in libro.sheetnames
    assert "Glosario" in libro.sheetnames
    hoja = libro["Plan NNN completo"]
    encabezados = [celda.value for celda in next(hoja.iter_rows())]
    assert "Nota" in encabezados


def test_word_estudiantil_es_abrible_y_declara_uso_educativo() -> None:
    salida = generar_word(_plan_minimo(), {"Tipo de paciente": "Adulto"})
    documento = Document(BytesIO(salida.getvalue()))
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)

    assert "finalidad educativa" in texto
    assert "requieren validación clínica" in texto


def test_word_docente_no_convierte_senales_en_calificacion() -> None:
    decisiones = {
        "Deterioro del intercambio gaseoso": {
            "decision": "Sin decidir",
            "jerarquia": "Complementario",
            "confianza": "Media",
            "criterios": [],
            "justificacion": "",
        }
    }
    analisis = analizar_sesion(decisiones, {})
    salida = generar_word_docente(_plan_minimo(), {}, decisiones, analisis)
    documento = Document(BytesIO(salida.getvalue()))
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)

    assert "NO dispone de una clave de respuesta validada" in texto
    assert "no califica decisiones como correctas o incorrectas" in texto
